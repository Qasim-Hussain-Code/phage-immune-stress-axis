"""Frozen source-data extraction and limited exploratory immune contrasts.

The workbook is read only. Source identifiers are retained literally and are
scoped by experiment and group. These identifiers do not constitute a virome
crosswalk. Missing outcomes are never imputed. Only the three endpoints and two
contrasts recorded in the feasibility gate are tested, as one six-test family.
"""

from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean, median, stdev, variance
from typing import Iterable, Sequence


SHEETS = {
    "Stress virome study I": ("study_i", {"Control": "ctr", "Stress": "ctr_stress"}),
    "FVT study II": (
        "study_ii",
        {"ctr": "ctr", "ctr_stress": "ctr_stress", "fvt_stress": "fvt_stress"},
    ),
}
ENDPOINTS = (
    "ly6chi_monocytes",
    "ly6chi_monocytes_cd62",
    "tnfa_splenocyte_unstimulated",
)
ENDPOINT_UNITS = {'ly6chi_monocytes':'percent of parent population',
                  'ly6chi_monocytes_cd62':'percent of parent population',
                  'tnfa_splenocyte_unstimulated':'pg/mL'}
EFFECT_UNITS = {'ly6chi_monocytes':'percentage points',
                'ly6chi_monocytes_cd62':'percentage points',
                'tnfa_splenocyte_unstimulated':'pg/mL'}
CONTRASTS = (("ctr_stress", "ctr"), ("fvt_stress", "ctr_stress"))
LONG_FIELDS = (
    "experiment_id", "source_sheet", "source_row", "source_column", "source_cell",
    "source_id", "mouse_id", "source_group", "group", "endpoint", "assay_domain",
    "selected_endpoint", "value", "observation_status",
)
INFERENTIAL_LIMITS = (
    "exploratory unadjusted animal-level comparison; cage allocation unavailable; "
    "upstream censoring and exclusions unresolved; no viral-immune association"
)


def _write_csv(path: Path, rows: list[dict], fields: Sequence[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        if not rows:
            raise ValueError(f"Refusing an empty scientific result table: {path}")
        fields = tuple(rows[0])
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _finite_values(values: Iterable[float]) -> list[float]:
    result = [float(value) for value in values]
    if any(not math.isfinite(value) for value in result):
        raise ValueError("Statistical inputs must contain finite observed values only")
    return result


def composite_mouse_id(experiment: str, group: str, source_id: object) -> str:
    """Keep punctuation and the literal Python string of the source cell value."""
    if source_id is None or str(source_id).strip() == "":
        raise ValueError("An individual observation requires an explicit source ID")
    identifier = str(source_id)
    if any("|" in part for part in (experiment, group, identifier)):
        raise ValueError("Identifiers may not contain the composite-key separator")
    return "|".join((experiment, group, identifier))


def _domain(endpoint: str) -> str:
    if endpoint.startswith(("sit_", "epm_", "fst_")):
        return "behavior"
    if "corticosterone" in endpoint or endpoint.startswith("cort_"):
        return "endocrine"
    return "immune"


def _measurement(value: object, location: str) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (float, int)):
        raise ValueError(f"Unsupported nonnumeric source value at {location}: {value!r}")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"Nonfinite source value at {location}")
    return result


def parse_workbook(path: Path) -> tuple[list[dict], list[dict]]:
    """Return all source phenotypes, including missing cells, and sheet inventory.

    Fully blank formatted rows are excluded. A partially observed animal is
    retained for every phenotype, with an explicit missing status where needed.
    Loading formulas rather than cached values makes unexpected formulas fail
    numeric validation instead of silently becoming missing measurements.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path, read_only=True, data_only=False)
    long_rows: list[dict] = []
    inventory: list[dict] = []
    try:
        if set(workbook.sheetnames) != set(SHEETS):
            raise ValueError(f"Unexpected source sheets: {workbook.sheetnames!r}")
        for sheet_name, (experiment, group_map) in SHEETS.items():
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headers = tuple(next(iterator))
            if headers[:2] != ("id", "group") or len(set(headers)) != len(headers):
                raise ValueError(f"Invalid or repeated headers in {sheet_name}")
            if any(not isinstance(name, str) or not name for name in headers):
                raise ValueError(f"Blank or nontext header in {sheet_name}")
            if experiment == "study_ii" and not set(ENDPOINTS).issubset(headers):
                raise ValueError("A prespecified study II endpoint is absent")
            group_counts: Counter = Counter()
            seen_keys: set[str] = set()
            blank_rows = 0
            observed_cells = 0
            for row_number, values in enumerate(iterator, 2):
                if all(value is None or value == "" for value in values):
                    blank_rows += 1
                    continue
                raw_id, raw_group = values[:2]
                if raw_group not in group_map:
                    raise ValueError(f"Unknown group at {sheet_name}!B{row_number}: {raw_group!r}")
                group = group_map[raw_group]
                mouse_id = composite_mouse_id(experiment, group, raw_id)
                if mouse_id in seen_keys:
                    raise ValueError(f"Duplicate biological identifier: {mouse_id}")
                seen_keys.add(mouse_id)
                group_counts[group] += 1
                for column_number, endpoint in enumerate(headers[2:], 3):
                    cell = f"{get_column_letter(column_number)}{row_number}"
                    value = _measurement(values[column_number - 1], f"{sheet_name}!{cell}")
                    observed_cells += value is not None
                    long_rows.append({
                        "experiment_id": experiment,
                        "source_sheet": sheet_name,
                        "source_row": row_number,
                        "source_column": column_number,
                        "source_cell": cell,
                        "source_id": str(raw_id),
                        "mouse_id": mouse_id,
                        "source_group": raw_group,
                        "group": group,
                        "endpoint": endpoint,
                        "assay_domain": _domain(endpoint),
                        "selected_endpoint": experiment == "study_ii" and endpoint in ENDPOINTS,
                        "value": "" if value is None else value,
                        "observation_status": "missing" if value is None else "observed",
                    })
            inventory.append({
                "experiment_id": experiment,
                "source_sheet": sheet_name,
                "source_max_row": sheet.max_row,
                "source_column_count": len(headers),
                "phenotype_column_count": len(headers) - 2,
                "nonempty_animal_rows": len(seen_keys),
                "fully_blank_rows_excluded": blank_rows,
                "observed_phenotype_cells": observed_cells,
                "missing_phenotype_cells": len(seen_keys) * (len(headers) - 2) - observed_cells,
                "group_counts_json": json.dumps(dict(sorted(group_counts.items())), separators=(",", ":")),
                "identifier_scope": "experiment|group|literal_source_id",
                "virome_linkage_verified": False,
            })
    finally:
        workbook.close()
    return long_rows, inventory


def prepare(root: Path, config: dict) -> dict:
    """Extract the unchanged source workbook and report phenotype missingness."""
    del config  # Endpoint and contrast choices are fixed by the feasibility gate.
    root = Path(root)
    long_rows, inventory = parse_workbook(root / "data/raw/source_data.xlsx")
    groups: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for row in long_rows:
        groups[(row["experiment_id"], row["group"], row["endpoint"])].append(row)
    missingness: list[dict] = []
    for (experiment, group, endpoint), rows in sorted(groups.items()):
        missing_ids = [row["source_id"] for row in rows if row["observation_status"] == "missing"]
        missingness.append({
            "experiment_id": experiment,
            "group": group,
            "endpoint": endpoint,
            "assay_domain": rows[0]["assay_domain"],
            "selected_endpoint": experiment == "study_ii" and endpoint in ENDPOINTS,
            "n_source_animals": len(rows),
            "n_observed": len(rows) - len(missing_ids),
            "n_missing": len(missing_ids),
            "missing_source_ids_json": json.dumps(missing_ids, separators=(",", ":")),
            "missing_reason": "unresolved upstream censoring/exclusion or unavailable assay" if missing_ids else "",
        })
    _write_csv(root / "data/processed/immune_long.csv", long_rows, LONG_FIELDS)
    _write_csv(root / "results/tables/immune_missingness.csv", missingness)
    _write_csv(root / "results/tables/source_inventory.csv", inventory)
    return {
        "status": "completed",
        "source_animal_rows": sum(row["nonempty_animal_rows"] for row in inventory),
        "source_animal_rows_by_experiment": {row["experiment_id"]: row["nonempty_animal_rows"] for row in inventory},
        "long_records": len(long_rows),
        "missing_records": sum(row["observation_status"] == "missing" for row in long_rows),
        "matched_processed_virome_profiles": 0,
        "outputs": ["data/processed/immune_long.csv", "results/tables/immune_missingness.csv", "results/tables/source_inventory.csv"],
    }


def welch_difference(first: Iterable[float], second: Iterable[float], alpha: float = 0.05) -> dict:
    """Unpaired mean(first)-mean(second), Welch interval and two-sided P value."""
    from scipy.stats import t as student_t

    if not 0 < alpha < 1:
        raise ValueError("alpha must lie strictly between zero and one")
    a, b = _finite_values(first), _finite_values(second)
    if len(a) < 2 or len(b) < 2:
        raise ValueError("Welch inference requires at least two observed animals per group")
    difference = mean(a) - mean(b)
    va, vb = variance(a) / len(a), variance(b) / len(b)
    se = math.sqrt(va + vb)
    if se == 0:
        raise ValueError("Welch inference is undefined when both sample variances are zero")
    degrees = (va + vb) ** 2 / (va ** 2 / (len(a) - 1) + vb ** 2 / (len(b) - 1))
    statistic = difference / se
    critical = float(student_t.ppf(1 - alpha / 2, degrees))
    return {
        "n_first": len(a),
        "n_second": len(b),
        "mean_first": mean(a),
        "mean_second": mean(b),
        "mean_difference": difference,
        "standard_error": se,
        "ci_low": difference - critical * se,
        "ci_high": difference + critical * se,
        "confidence_level": 1 - alpha,
        "t_statistic": statistic,
        "degrees_freedom": degrees,
        "p_value": float(2 * student_t.sf(abs(statistic), degrees)),
    }


def benjamini_hochberg(p_values: Sequence[float]) -> list[float]:
    """BH adjustment across the complete fixed family, preserving input order."""
    p = _finite_values(p_values)
    if any(value < 0 or value > 1 for value in p):
        raise ValueError("P values must lie in [0,1]")
    order = sorted(range(len(p)), key=lambda index: (p[index], index))
    adjusted = [0.0] * len(p)
    running = 1.0
    for rank_index in range(len(order) - 1, -1, -1):
        index = order[rank_index]
        running = min(running, p[index] * len(p) / (rank_index + 1))
        adjusted[index] = running
    return adjusted


def leave_one_out(first: Sequence[tuple[str, float]], second: Sequence[tuple[str, float]], alpha: float = 0.05) -> dict:
    """Recompute each contrast after dropping every observed animal once."""
    base = welch_difference([x[1] for x in first], [x[1] for x in second], alpha)
    trials: list[tuple[str, dict]] = []
    for arm, group in enumerate((first, second)):
        for index, (identifier, _) in enumerate(group):
            a = [value for j, (_, value) in enumerate(first) if arm != 0 or j != index]
            b = [value for j, (_, value) in enumerate(second) if arm != 1 or j != index]
            trials.append((identifier, welch_difference(a, b, alpha)))
    low = min(trials, key=lambda entry: entry[1]["mean_difference"])
    high = max(trials, key=lambda entry: entry[1]["mean_difference"])
    return {
        "loo_n_refits": len(trials),
        "loo_mean_difference_min": low[1]["mean_difference"],
        "loo_mean_difference_max": high[1]["mean_difference"],
        "loo_min_difference_omitted_mouse": low[0],
        "loo_max_difference_omitted_mouse": high[0],
        "loo_p_value_min": min(row["p_value"] for _, row in trials),
        "loo_p_value_max": max(row["p_value"] for _, row in trials),
        "loo_direction_preserved": all(row["mean_difference"] * base["mean_difference"] > 0 for _, row in trials),
    }


def _read_long(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if tuple(reader.fieldnames or ()) != LONG_FIELDS:
            raise ValueError("Unexpected prepared phenotype schema; rerun preparation")
        rows = list(reader)
    seen: set[tuple[str, str]] = set()
    for row in rows:
        expected = composite_mouse_id(row["experiment_id"], row["group"], row["source_id"])
        if row["mouse_id"] != expected:
            raise ValueError("Prepared source identifier has changed")
        key = (row["mouse_id"], row["endpoint"])
        if key in seen:
            raise ValueError(f"Repeated individual endpoint record: {key}")
        seen.add(key)
        status = row["observation_status"]
        if status not in ("observed", "missing") or ((status == "missing") != (row["value"] == "")):
            raise ValueError(f"Inconsistent missingness for {key}")
        if status == "observed":
            _finite_values([float(row["value"])])
    return rows


def analyze(root: Path, config: dict) -> dict:
    """Analyze the fixed study II family separately from unavailable viral data."""
    from scipy.stats import mannwhitneyu

    del config
    root = Path(root)
    rows = _read_long(root / "data/processed/immune_long.csv")
    selected = [row for row in rows if row["experiment_id"] == "study_ii" and row["endpoint"] in ENDPOINTS]
    data: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in selected:
        data[(row["endpoint"], row["group"])].append(row)
    summary: list[dict] = []
    for endpoint in ENDPOINTS:
        for group in ("ctr", "ctr_stress", "fvt_stress"):
            observations = data[(endpoint, group)]
            values = [float(row["value"]) for row in observations if row["observation_status"] == "observed"]
            if len(values) < 2:
                raise ValueError(f"Insufficient observations for selected endpoint: {endpoint}, {group}")
            summary.append({
                "experiment_id": "study_ii",
                "endpoint": endpoint,
                "group": group,
                "n_source_animals": len(observations),
                "n_observed": len(values),
                "n_missing": len(observations) - len(values),
                "mean": mean(values),
                "standard_deviation": stdev(values),
                "median": median(values),
                "minimum": min(values),
                "maximum": max(values),
                "scale": "untransformed source units",
                "units": ENDPOINT_UNITS[endpoint],
            })
    contrasts: list[dict] = []
    sensitivities: list[dict] = []
    for endpoint in ENDPOINTS:
        for first_group, second_group in CONTRASTS:
            first = [(row["mouse_id"], float(row["value"])) for row in data[(endpoint, first_group)] if row["observation_status"] == "observed"]
            second = [(row["mouse_id"], float(row["value"])) for row in data[(endpoint, second_group)] if row["observation_status"] == "observed"]
            label = f"{first_group}_minus_{second_group}"
            stats = welch_difference([x[1] for x in first], [x[1] for x in second])
            contrasts.append({
                "experiment_id": "study_ii",
                "endpoint": endpoint,
                "contrast": label,
                "first_group": first_group,
                "second_group": second_group,
                **stats,
                "method": "two-sided unpaired Welch t-test",
                "scale": "untransformed source units",
                "effect_units": EFFECT_UNITS[endpoint],
                "family": "three fixed endpoints by two fixed contrasts",
                "family_size": len(ENDPOINTS) * len(CONTRASTS),
                "interpretation": INFERENTIAL_LIMITS,
            })
            mw = mannwhitneyu([x[1] for x in first], [x[1] for x in second], alternative="two-sided", method="auto")
            sensitivities.append({
                "experiment_id": "study_ii",
                "endpoint": endpoint,
                "contrast": label,
                "mann_whitney_u_first": float(mw.statistic),
                "mann_whitney_p_value": float(mw.pvalue),
                "rank_method": "two-sided Mann-Whitney U; scipy method=auto; distributional sensitivity",
                **leave_one_out(first, second),
                "interpretation": "sensitivity diagnostics; not independent confirmatory tests; no additional exclusions",
            })
    if len(contrasts) != 6:
        raise AssertionError("The fixed inferential family must contain exactly six tests")
    q_values = benjamini_hochberg([row["p_value"] for row in contrasts])
    rank_q = benjamini_hochberg([row["mann_whitney_p_value"] for row in sensitivities])
    for row, q_value in zip(contrasts, q_values):
        row["q_value"] = q_value
    for row, q_value in zip(sensitivities, rank_q):
        row["mann_whitney_q_value_sensitivity_family"] = q_value
    _write_csv(root / "results/tables/immune_contrasts.csv", contrasts)
    _write_csv(root / "results/tables/immune_sensitivity.csv", sensitivities)
    _write_csv(root / "results/tables/immune_summary.csv", summary)
    return {
        "status": "completed",
        "analysis": "separate exploratory immune contrasts",
        "tests": len(contrasts),
        "bh_family_size": 6,
        "n_q_below_0_05": sum(row["q_value"] < 0.05 for row in contrasts),
        "selected_endpoints": list(ENDPOINTS),
        "contrast_results": contrasts,
        "interpretation": INFERENTIAL_LIMITS,
        "outputs": ["results/tables/immune_contrasts.csv", "results/tables/immune_sensitivity.csv", "results/tables/immune_summary.csv"],
    }


def run(root: Path, config: dict) -> dict:
    """Convenience entry point; numbered stages normally call prepare/analyze."""
    return {"prepare": prepare(root, config), "analyze": analyze(root, config)}
