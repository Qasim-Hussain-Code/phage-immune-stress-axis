"""Conservative specimen inventory and source-local assay linkage.

Accession metadata is evidence that a specimen was sequenced. It is not an
abundance profile, nor a crosswalk to the source workbook. This module never
merges those two sources on a guessed animal identifier.
"""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable


LINKAGE_FIELDS = [
    "sample_id", "mouse_id", "pool_id", "experiment_id", "group", "timepoint",
    "virome_available", "bacteriome_available", "immune_available",
    "behavior_available", "matched_assays", "linkage_confidence", "sample_type",
    "source", "run_id", "metadata_candidate_mouse_id", "note",
    "raw_reads_deposited",
]
INVENTORY_FIELDS = [
    "accession", "sample_type", "run_count", "sample_count", "fastq_file_count",
    "compressed_fastq_bytes", "compressed_fastq_gb_decimal", "source",
    "analysis_role",
]
SHEETS = {
    "Stress virome study I": "stress_study_i",
    "FVT study II": "fvt_study_ii",
}
BEHAVIOR_PREFIXES = ("sit_", "fst_", "epm_")
IMMUNE_PREFIXES = (
    "il6_", "il10_", "il12_", "tnfa_", "ifny_", "ip10_", "ly6chi_",
    "mhc2_", "neutrophils_", "cd62_", "thelper_", "cd44hi_", "cd69_",
)


def literal_id(value: Any) -> str:
    """Preserve a source identifier without punctuation removal or rounding."""
    if value is None or isinstance(value, bool):
        raise ValueError("A nonempty source identifier is required")
    result = str(value)
    if not result or not result.strip():
        raise ValueError("A nonempty source identifier is required")
    return result


def source_animal_key(experiment: str, group: str, animal_id: Any) -> str:
    """Construct a source-local key; repeated numeric IDs are not global IDs."""
    if not experiment or not group:
        raise ValueError("Animal identifiers require experiment and group")
    # JSON escaping avoids collisions when a label itself contains separators.
    return json.dumps([experiment, group, literal_id(animal_id)], ensure_ascii=False)


def _empty_record() -> dict[str, Any]:
    row = dict.fromkeys(LINKAGE_FIELDS, "")
    for name in (
        "virome_available", "bacteriome_available", "immune_available",
        "behavior_available", "raw_reads_deposited",
    ):
        row[name] = False
    return row


def _workbook_records(
    rows: Iterable[tuple[Any, ...]], experiment: str, sheet: str,
    source: str = "data/raw/source_data.xlsx",
) -> list[dict[str, Any]]:
    """Read observed assays while leaving workbook animal IDs source-local."""
    rows = iter(rows)
    try:
        header = next(rows)
    except StopIteration as exc:
        raise ValueError(f"Empty source sheet: {sheet}") from exc
    names = [str(value) if value is not None else "" for value in header]
    if names[:2] != ["id", "group"]:
        raise ValueError(f"Unexpected source workbook header in {sheet}")
    records: list[dict[str, Any]] = []
    seen: set[str] = set()
    for excel_row, values in enumerate(rows, 2):
        if not any(value is not None for value in values):
            continue
        if len(values) < 2 or values[0] is None or values[1] is None:
            raise ValueError(f"Missing animal ID or group at {sheet}!{excel_row}")
        animal, group = literal_id(values[0]), str(values[1])
        key = source_animal_key(experiment, group, animal)
        if key in seen:
            raise ValueError(f"Duplicate source animal key at {sheet}!{excel_row}: {key}")
        seen.add(key)
        observed = [
            name for name, value in zip(names, values)
            if value is not None and name not in ("id", "group", "")
        ]
        immune = any(name.startswith(IMMUNE_PREFIXES) for name in observed)
        behavior = any(name.startswith(BEHAVIOR_PREFIXES) for name in observed)
        row = _empty_record()
        row.update({
            "sample_id": "workbook:" + key,
            "mouse_id": animal,
            "experiment_id": experiment,
            "group": group,
            "timepoint": "endpoint_and_FST_assays",
            "immune_available": immune,
            "behavior_available": behavior,
            "matched_assays": "immune;behavior" if immune and behavior else "",
            "linkage_confidence": "within_workbook_row_only",
            "sample_type": "individual_animal_source_row",
            "source": f"{source}|sheet={sheet}|row={excel_row}",
            "note": (
                "Constructed sample_id namespaces the literal workbook ID by experiment and group. "
                "Same-row observed immune/behavior columns are linked; no sequence or abundance "
                "profile is linked. Empty cells remain missing. Timepoint is an assay-scope "
                "description, not a source-provided sampling date."
            ),
        })
        records.append(row)
    return records


def classify_sequencing_sample(alias: str, attributes: dict[str, str]) -> str:
    """Use explicit pool/gavage labels, never infer individual pool members."""
    if "pool" in alias.lower():
        return "pooled_fecal_virome"
    if attributes.get("gavage_status", "").lower() == "gavage" or "gavage" in alias.lower():
        return "transplant_inoculum"
    return "individual_recipient_bulk_metagenome"


def _sequencing_record(run_rows: list[dict[str, str]], attributes: dict[str, str]) -> dict[str, Any]:
    run = run_rows[0]
    sample_accession, alias = run["sample_accession"], run["sample_alias"]
    if alias != attributes.get("sample_alias"):
        raise ValueError(f"Run/BioSample alias disagreement for {sample_accession}")
    sample_type = classify_sequencing_sample(alias, attributes)
    group = attributes.get("sample_type", "")
    replicate = attributes.get("replicate_num", "")
    if replicate.strip().lower() in ("na", "n/a", "none", ""):
        replicate = ""
    row = _empty_record()
    row.update({
        "sample_id": sample_accession,
        "experiment_id": "fvt_study_ii",
        "group": group,
        "timepoint": attributes.get("gavage_status", ""),
        "sample_type": sample_type,
        "source": "data/raw/ritz2024_archive/PRJNA970614_sample_attributes.csv",
        "run_id": ";".join(r["run_accession"] for r in run_rows),
        "raw_reads_deposited": True,
        "linkage_confidence": "accession_metadata_only_no_assay_crosswalk",
    })
    if sample_type == "pooled_fecal_virome":
        row["pool_id"] = alias
        row["note"] = (
            "Literal pool alias retained. Pool membership and repeated pool composition "
            "are unresolved. No individual mouse or immune outcome assigned."
        )
    elif sample_type == "transplant_inoculum":
        row["note"] = (
            "Gavage/inoculum specimen. replicate_num is not treated as a recipient "
            "intestinal measurement. No inoculum-to-animal phenotype crosswalk verified."
        )
    else:
        if replicate and group:
            row["metadata_candidate_mouse_id"] = source_animal_key(
                "PRJNA970614", group, replicate,
            )
        row["note"] = (
            f"Literal sequencing alias={alias}; replicate_num={replicate or 'missing'}. "
            "Before/After uses explicit BioSample gavage_status. Candidate ID is "
            "source-local metadata, not a workbook crosswalk. Bulk reads were deposited; "
            "processed viral/bacterial abundance matrices were not retrieved."
        )
    return row


def _read_csv(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def _fastq_counts(run: dict[str, str]) -> tuple[int, int]:
    fields = [entry.strip() for entry in run.get("fastq_bytes", "").split(";") if entry.strip()]
    values = [int(entry) for entry in fields]
    if any(value < 0 for value in values):
        raise ValueError(f"Negative FASTQ size for {run.get('run_accession')}")
    if not values:
        raise ValueError(f"Missing FASTQ size for {run.get('run_accession')}")
    return len(values), sum(values)


def _inventory_row(accession: str, category: str, runs: list[dict[str, str]], role: str) -> dict[str, Any]:
    sizes = [_fastq_counts(run) for run in runs]
    total_bytes = sum(size for _, size in sizes)
    return {
        "accession": accession,
        "sample_type": category,
        "run_count": len(runs),
        "sample_count": len({run["sample_accession"] for run in runs}),
        "fastq_file_count": sum(count for count, _ in sizes),
        "compressed_fastq_bytes": total_bytes,
        "compressed_fastq_gb_decimal": f"{total_bytes / 1_000_000_000:.6f}",
        "source": f"data/raw/ritz2024_archive/{accession}_ena_runs.tsv",
        "analysis_role": role,
    }


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def audit(root: Path | str, config: dict[str, Any]) -> dict[str, Any]:
    """Build the validated availability audit from frozen small public inputs.

    ``config`` is accepted for the common stage interface. Identifiers and
    inclusion are determined by the frozen sources rather than configuration.
    No network access, raw-read download or cross-source animal merge occurs.
    """
    root = Path(root)
    archive = root / "data/raw/ritz2024_archive"
    runs = _read_csv(archive / "PRJNA970614_ena_runs.tsv", "\t")
    attributes = _read_csv(archive / "PRJNA970614_sample_attributes.csv")
    by_accession: dict[str, dict[str, str]] = {}
    for row in attributes:
        accession = row["sample_accession"]
        if accession in by_accession:
            raise ValueError(f"Duplicate BioSample attributes: {accession}")
        by_accession[accession] = row
    if len({run["run_accession"] for run in runs}) != len(runs):
        raise ValueError("Duplicate sequencing run accessions")
    specimen_runs: dict[str, list[dict[str, str]]] = defaultdict(list)
    categories: dict[str, list[dict[str, str]]] = defaultdict(list)
    for run in runs:
        accession = run["sample_accession"]
        if accession not in by_accession:
            raise ValueError(f"Missing BioSample attributes for {accession}")
        specimen_runs[accession].append(run)
        categories[classify_sequencing_sample(run["sample_alias"], by_accession[accession])].append(run)
    sequence_rows = [
        _sequencing_record(specimen_runs[accession], by_accession[accession])
        for accession in sorted(specimen_runs)
    ]

    # openpyxl is imported only when the audit needs the source XLSX.
    from openpyxl import load_workbook
    workbook = load_workbook(root / "data/raw/source_data.xlsx", read_only=True, data_only=True)
    source_rows: list[dict[str, Any]] = []
    try:
        for sheet, experiment in SHEETS.items():
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Missing source-data worksheet: {sheet}")
            source_rows.extend(_workbook_records(
                workbook[sheet].iter_rows(values_only=True), experiment, sheet,
            ))
    finally:
        workbook.close()
    linkage = sequence_rows + source_rows
    if len({row["sample_id"] for row in linkage}) != len(linkage):
        raise ValueError("Audit produced duplicate specimen/source-row keys")

    inventory = [
        _inventory_row("PRJNA970614", category, category_runs, "raw_reads_only_no_processed_profile")
        for category, category_runs in sorted(categories.items())
    ]
    for accession, category, role in [
        ("PRJEB69687", "16S_amplicon_manifest_only", "initial_experiment_raw_manifest_no_animal_crosswalk"),
        ("PRJEB67706", "brain_transcriptomic_manifest_only", "secondary_assay_outside_selected_analysis"),
    ]:
        extra_runs = _read_csv(archive / f"{accession}_ena_runs.tsv", "\t")
        inventory.append(_inventory_row(accession, category, extra_runs, role))

    group_counts = Counter((row["experiment_id"], row["group"]) for row in source_rows)
    candidates: dict[str, set[str]] = defaultdict(set)
    for row in sequence_rows:
        if row["metadata_candidate_mouse_id"]:
            candidates[row["metadata_candidate_mouse_id"]].add(row["timepoint"])
    summary = {
        "principal_accession": "PRJNA970614",
        "principal_run_count": len(runs),
        "principal_specimen_count": len(sequence_rows),
        "sequencing_category_counts": {category: len(rows) for category, rows in sorted(categories.items())},
        "principal_compressed_fastq_bytes": sum(_fastq_counts(run)[1] for run in runs),
        "workbook_animal_rows": len(source_rows),
        "workbook_group_counts": [
            {"experiment_id": experiment, "group": group, "animal_rows": count}
            for (experiment, group), count in sorted(group_counts.items())
        ],
        "workbook_rows_with_observed_immune": sum(bool(row["immune_available"]) for row in source_rows),
        "workbook_rows_with_observed_behavior": sum(bool(row["behavior_available"]) for row in source_rows),
        "within_workbook_immune_behavior_rows": sum(bool(row["matched_assays"]) for row in source_rows),
        "metadata_candidate_animal_keys": len(candidates),
        "metadata_candidates_with_before_after": sum({"Before", "After"} <= times for times in candidates.values()),
        "processed_virome_profiles_retrieved": 0,
        "verified_virome_immune_matches": 0,
        "verified_bacteriome_immune_matches": 0,
        "longitudinal_phageome_stability_measurable": False,
        "primary_association_identifiable": False,
        "linkage_rows": len(linkage),
        "additional_accessions_are_inventory_only": ["PRJEB69687", "PRJEB67706"],
        "decision": "Separate source-row immune analysis and sequencing availability audit only.",
    }
    tables = root / "results/tables"
    reports = root / "results/reports"
    tables.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)
    _write_csv(tables / "sample_linkage_audit.csv", LINKAGE_FIELDS, linkage)
    _write_csv(tables / "sequencing_inventory.csv", INVENTORY_FIELDS, inventory)
    (reports / "sample_audit_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8",
    )
    count_text = ", ".join(f"{category}: {len(rows)}" for category, rows in sorted(categories.items()))
    group_text = "\n".join(
        f"| {experiment} | {group} | {count} |"
        for (experiment, group), count in sorted(group_counts.items())
    )
    report = f"""# Sample linkage audit

This report is generated from frozen accession manifests, BioSample attributes,
and the supplementary source workbook. No raw reads were downloaded.

## Counts and source scope

PRJNA970614 contains {len(runs)} runs representing {len(sequence_rows)} sequencing
specimens. Categories from explicit aliases/BioSample attributes are {count_text}.
These specimen counts are not independent animal counts. Multiple runs per
specimen, if encountered, share an accession row and remain technical sequencing
observations. Compressed FASTQ sizes total
{summary['principal_compressed_fastq_bytes']:,} bytes before processing.

The source workbook contributes {len(source_rows)} animal rows:

| Experiment | Literal group | Animal rows |
|---|---|---:|
{group_text}

The linkage CSV includes these principal sequencing specimens and source-workbook
animal rows. PRJEB69687 and PRJEB67706 are represented by totals in
`sequencing_inventory.csv`; their aliases are not assigned to workbook animals.

## Identifier policy

`sample_id` is a literal BioSample accession for a sequencing specimen. Workbook
keys are explicitly constructed from experiment, literal group, and literal ID.
They identify a source row and do not assert a shared upstream animal naming
scheme. Workbook `mouse_id` remains source-local. Initial-study IDs restart in
each treatment group. Decimal identifiers such as `1.1`, `1.2`, `19.1`, and `19.2`
retain their punctuation and never become integer IDs. No match uses row order,
punctuation removal, numeric rounding, or presumed renumbering.

Sequencing `replicate_num` values are retained only as experiment/accession and
group-qualified `metadata_candidate_mouse_id` keys. They do not populate
`mouse_id`, and no candidate is joined to the workbook. This metadata contains
{len(candidates)} candidate animal keys, with
{summary['metadata_candidates_with_before_after']} having both explicit Before
and After specimen annotations. These are candidate repeated raw specimens,
not calculated longitudinal viral profiles. Before/After comes from the
BioSample `gavage_status` field. Pool aliases populate `pool_id`; pools have no
individual mouse ID or candidate ID. Inocula remain exposure material and are
never recipient intestinal observations. Heat-inactivated FVT specimens remain
in the sequencing inventory despite no corresponding source-workbook group.

## What is available and linked

The boolean `virome_available` and `bacteriome_available` fields mean a retrieved
processed abundance profile, not the existence of deposited raw reads. Both are
false for all audit rows. `raw_reads_deposited` distinguishes sequenced specimens.
The workbook has {summary['workbook_rows_with_observed_immune']} rows with at
least one observed immune measurement and
{summary['workbook_rows_with_observed_behavior']} with a behavioral measurement.
There are {summary['within_workbook_immune_behavior_rows']} same-row immune and
behavior links. Individual endpoint missingness is not summarized by these
any-assay flags. Corticosterone is treated as endocrine and does not create an
immune-available flag. The source's late incomplete animal is retained.

No retrieved abundance profile is linked to an immune outcome. The original
paper describes individual bulk-metagenomic reads mapped to a gavage-derived
viral catalog, alongside pooled enriched viromes. Therefore the obstacle is
missing processed inputs and an unverified crosswalk, rather than a claim that
all original viral measurements were pooled. Do not infer phage-immune
associations, bacterial adjustment, recipient viral load from inoculum spikes,
or measured longitudinal stability from this inventory.

Source-row immune/behavior links remain valid within their sheet. Whole-study
mouse and pool membership across assays remain unresolved. See
`data_feasibility.md` and `docs/paper_audit_notes.md` for experimental details.
"""
    (reports / "sample_linkage.md").write_text(report, encoding="utf-8")
    return summary
