"""Independent checks for source extraction and the fixed immune estimators.

All created workbook values below are synthetic software fixtures, never
biological results. Run with unittest discovery from the repository root.
"""

from __future__ import annotations

import csv
import math
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from pisa.immune import (
    ENDPOINTS, LONG_FIELDS, SHEETS, analyze, benjamini_hochberg,
    composite_mouse_id, leave_one_out, parse_workbook, prepare, welch_difference,
)


class TestSourceIdentifiers(unittest.TestCase):
    def test_literal_and_composite_ids(self):
        self.assertEqual(composite_mouse_id("study_ii", "ctr", 19.2), "study_ii|ctr|19.2")
        self.assertEqual(composite_mouse_id("study_ii", "ctr", "01.20"), "study_ii|ctr|01.20")
        self.assertEqual(composite_mouse_id("study_i", "ctr", 1), "study_i|ctr|1")
        self.assertNotEqual(composite_mouse_id("study_i", "ctr", 1), composite_mouse_id("study_i", "ctr_stress", 1))
        self.assertNotEqual(composite_mouse_id("study_i", "ctr", 1), composite_mouse_id("study_ii", "ctr", 1))
        self.assertNotEqual(composite_mouse_id("study_ii", "ctr", 1.2), composite_mouse_id("study_ii", "ctr", 12))
        with self.assertRaises(ValueError):
            composite_mouse_id("study_ii", "ctr", None)
        with self.assertRaises(ValueError):
            composite_mouse_id("study_ii", "ctr", "a|b")

    def fixture(self, root: Path, duplicate=False, unknown_group=False) -> Path:
        from openpyxl import Workbook

        path = root / "data/raw/source_data.xlsx"
        path.parent.mkdir(parents=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        first = workbook.create_sheet("Stress virome study I")
        first.append(["id", "group", "il6_cona"])
        first.append([1, "Control", 0])
        first.append([1, "Stress", None])
        sheet = workbook.create_sheet("FVT study II")
        sheet.append(["id", "group", *ENDPOINTS])
        for group in ("ctr", "ctr_stress", "fvt_stress"):
            for index in range(4):
                key = "19.2" if index == 0 else str(index)
                values = [index + 1.0, index + 3.0, index + 8.0]
                if index == 1:
                    values[2] = None
                sheet.append([key, "unknown" if unknown_group else group, *values])
        if duplicate:
            sheet.append(["19.2", "ctr", 1, 2, 3])
        # A physically present but completely blank row must not become a mouse.
        sheet.cell(40, 1).number_format = "0.0"
        workbook.save(path)
        workbook.close()
        return path

    def test_blank_rows_missing_outcomes_and_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self.fixture(root)
            before = source.read_bytes()
            rows, inventory = parse_workbook(source)
            self.assertEqual(source.read_bytes(), before)
            self.assertEqual([x["nonempty_animal_rows"] for x in inventory], [2, 12])
            self.assertEqual(inventory[1]["fully_blank_rows_excluded"], 27)
            zero = next(row for row in rows if row["experiment_id"] == "study_i" and row["group"] == "ctr")
            self.assertEqual(zero["value"], 0)
            self.assertEqual(zero["observation_status"], "observed")
            missing = [row for row in rows if row["observation_status"] == "missing"]
            self.assertEqual(len(missing), 4)
            self.assertTrue(all(row["value"] == "" for row in missing))
            observed_key = next(row for row in rows if row["source_id"] == "19.2")
            self.assertEqual(observed_key["mouse_id"], "study_ii|ctr|19.2")
            self.assertEqual(observed_key["source_cell"], "C2")
            prepared = prepare(root, {})
            self.assertEqual(prepared["source_animal_rows"], 14)
            self.assertEqual(prepared["matched_processed_virome_profiles"], 0)
            with (root / "data/processed/immune_long.csv").open(newline="", encoding="utf-8") as handle:
                reader = csv.DictReader(handle)
                self.assertEqual(tuple(reader.fieldnames), LONG_FIELDS)
                saved = list(reader)
            self.assertEqual(len(saved), 38)
            self.assertEqual(sum(row["value"] == "" for row in saved), 4)

    def test_duplicate_and_unknown_groups_fail(self):
        for option in ("duplicate", "unknown_group"):
            with self.subTest(option=option), tempfile.TemporaryDirectory() as directory:
                path = self.fixture(Path(directory), **{option: True})
                with self.assertRaises(ValueError):
                    parse_workbook(path)

    def test_unavailable_source_does_not_make_results(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaises(FileNotFoundError):
                prepare(root, {})
            self.assertFalse((root / "data/processed/immune_long.csv").exists())

    def test_fixed_six_test_family_and_no_study_i_tests(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.fixture(root)
            prepare(root, {})
            result = analyze(root, {})
            self.assertEqual(result["tests"], 6)
            self.assertEqual(result["bh_family_size"], 6)
            self.assertTrue(all(row["experiment_id"] == "study_ii" for row in result["contrast_results"]))
            self.assertEqual({row["endpoint"] for row in result["contrast_results"]}, set(ENDPOINTS))
            self.assertEqual({row["n_first"] for row in result["contrast_results"] if row["endpoint"] == ENDPOINTS[2]}, {3})


class TestIndependentStatistics(unittest.TestCase):
    def test_welch_matches_scipy(self):
        from scipy.stats import ttest_ind

        examples = [
            ([1.0, 2.0, 3.0], [3.0, 4.0, 8.0, 12.0]),
            ([1, 1, 1, 1], [2, 3, 5]),
            ([0, 0.02, 0.5, 60, 80], [0.01, 0.03, 0.5]),
            ([1.3, 5.2, 7.8, 3.1, 6.4], [2.2, 4.1, 8.5, 9.7]),
        ]
        for first, second in examples:
            with self.subTest(first=first):
                actual = welch_difference(first, second)
                if len(set(first))==1 or len(set(second))==1:
                    # This warning is an expected property of the deliberately
                    # constant software fixture, not a biological-data warning.
                    with self.assertWarnsRegex(RuntimeWarning,'Precision loss'):
                        reference = ttest_ind(first, second, equal_var=False)
                else:
                    reference = ttest_ind(first, second, equal_var=False)
                interval = reference.confidence_interval()
                self.assertAlmostEqual(actual["t_statistic"], reference.statistic, places=12)
                self.assertAlmostEqual(actual["p_value"], reference.pvalue, places=12)
                self.assertAlmostEqual(actual["degrees_freedom"], reference.df, places=12)
                self.assertAlmostEqual(actual["ci_low"], interval.low, places=10)
                self.assertAlmostEqual(actual["ci_high"], interval.high, places=10)

    def test_bh_matches_statsmodels(self):
        from statsmodels.stats.multitest import multipletests

        for values in ([0.001, 0.005, 0.02, 0.04, 0.4, 0.9], [1, 0, 0.04, 0.04, 0.8, 0.01], [0.1] * 6):
            expected = multipletests(values, method="fdr_bh")[1]
            actual = benjamini_hochberg(values)
            self.assertTrue(all(abs(a - b) < 1e-14 for a, b in zip(actual, expected)))

    def test_analytical_identity_and_edge_cases(self):
        first, second = [1, 2, 3, 8], [2, 5, 6]
        actual = welch_difference(first, second)
        reversed_result = welch_difference(second, first)
        self.assertAlmostEqual(actual["mean_difference"], -reversed_result["mean_difference"])
        self.assertAlmostEqual(actual["ci_low"], -reversed_result["ci_high"])
        self.assertAlmostEqual(actual["p_value"], reversed_result["p_value"])
        for first, second in (([1], [2, 3]), ([1, math.nan], [2, 3]), ([1, 1], [2, 2])):
            with self.assertRaises(ValueError):
                welch_difference(first, second)
        with self.assertRaises(ValueError):
            benjamini_hochberg([0.1, math.nan])
        with self.assertRaises(ValueError):
            benjamini_hochberg([-0.1, 0.2])

    def test_leave_one_out_agrees_with_brute_force_means(self):
        first = [(f"a{index}", value) for index, value in enumerate([1, 2, 4, 80])]
        second = [(f"b{index}", value) for index, value in enumerate([3, 5, 6])]
        expected = []
        for side, values in enumerate((first, second)):
            for index in range(len(values)):
                a = [value for j, (_, value) in enumerate(first) if side != 0 or j != index]
                b = [value for j, (_, value) in enumerate(second) if side != 1 or j != index]
                expected.append(sum(a) / len(a) - sum(b) / len(b))
        actual = leave_one_out(first, second)
        self.assertEqual(actual["loo_n_refits"], 7)
        self.assertAlmostEqual(actual["loo_mean_difference_min"], min(expected))
        self.assertAlmostEqual(actual["loo_mean_difference_max"], max(expected))
        self.assertFalse(actual["loo_direction_preserved"])
        self.assertEqual(actual["loo_min_difference_omitted_mouse"], "a3")


if __name__ == "__main__":
    unittest.main()
